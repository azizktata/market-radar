"""
Excel export — generates a .xlsx file matching the spec column layout.
Lowest price per row is highlighted in green; OK/KO status based on PVC comparison.
"""

import io
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

GREEN_FILL  = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
RED_FILL    = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
HEADER_FILL = PatternFill(start_color="2D6A4F", end_color="2D6A4F", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def build_excel(products: list[dict], sites: list[dict]) -> bytes:
    """
    products: list of dicts returned by GET /api/products
    sites: list of enabled site dicts with keys: scraper_key, name, threshold
    """
    columns = ["Référence", "Nom", "Marque", "Catégorie", "Sous-catégorie", "PVC"]
    for s in sites:
        label = s["name"]
        columns += [f"Prix {label}", f"Statut {label}", f"URL {label}"]
    columns.append("Dernière mise à jour")

    wb = Workbook()
    ws = wb.active
    ws.title = "Market Radar"

    # Header row
    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "A2"

    # Data rows
    for row_idx, product in enumerate(products, start=2):
        prices: list[tuple[float, int]] = []  # (price, col_index) for highlighting

        ws.cell(row=row_idx, column=1, value=product.get("reference", ""))
        ws.cell(row=row_idx, column=2, value=product.get("name") or "")
        ws.cell(row=row_idx, column=3, value=product.get("marque") or "")
        ws.cell(row=row_idx, column=4, value=product.get("category") or "")
        ws.cell(row=row_idx, column=5, value=product.get("sous_categorie") or "")
        ws.cell(row=row_idx, column=6, value=product.get("pvc") or "")

        col = 7
        last_scraped = None
        pvc = product.get("pvc")

        for s in sites:
            site_key = s["scraper_key"]
            threshold = s.get("threshold", 0)
            data = product.get(site_key) or {}
            price = data.get("price")
            price_raw = data.get("price_raw") or ""
            url = data.get("url") or ""
            scraped_at = data.get("scraped_at")

            # OK/KO status
            if pvc and price is not None:
                status_val = "OK" if price <= pvc * (1 + threshold / 100) else "KO"
            else:
                status_val = ""

            ws.cell(row=row_idx, column=col,     value=price_raw)
            ws.cell(row=row_idx, column=col + 1, value=status_val)
            link_cell = ws.cell(row=row_idx, column=col + 2, value=url)
            if url:
                link_cell.hyperlink = url
                link_cell.style = "Hyperlink"

            if price is not None:
                prices.append((price, col))

            if scraped_at and (last_scraped is None or scraped_at > last_scraped):
                last_scraped = scraped_at

            col += 3

        ws.cell(row=row_idx, column=len(columns), value=last_scraped or "")

        # Highlight cheapest / most expensive price in each row
        if len(prices) > 1:
            min_price = min(p for p, _ in prices)
            max_price = max(p for p, _ in prices)
            for price, col_idx in prices:
                cell = ws.cell(row=row_idx, column=col_idx)
                if price == min_price:
                    cell.fill = GREEN_FILL
                elif price == max_price:
                    cell.fill = RED_FILL

    # Auto-size columns
    for col_idx, col_name in enumerate(columns, start=1):
        max_len = len(col_name) + 2
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = min(max_len, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
